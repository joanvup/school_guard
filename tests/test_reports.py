import io
from datetime import datetime, time

from openpyxl import load_workbook

from app import models
from app.database import SessionLocal
from app.utils.time import today_local
from tests.conftest import USER_PASSWORD
from tests.helpers import create_student, get_door_id, login

TODAY = today_local().isoformat()


def _create_exit_logs(student_id, door_id, count=1, **snapshots):
    db = SessionLocal()
    try:
        operator = db.query(models.User).filter_by(username="operator").one()
        for _ in range(count):
            db.add(
                models.ExitLog(
                    student_id=student_id,
                    operator_id=operator.id,
                    door_id=door_id,
                    timestamp=datetime.combine(today_local(), time(10, 0)),
                    student_name=snapshots.get("student_name", "Ana Perez"),
                    course=snapshots.get("course", "10A"),
                    door_name=snapshots.get("door_name", "Puerta Principal"),
                    authorized_at_scan=True,
                )
            )
        db.commit()
    finally:
        db.close()


def test_reports_paginate_and_use_frozen_snapshots(client):
    login(client, "operator", USER_PASSWORD)
    student = create_student()
    _create_exit_logs(student.id, get_door_id(), count=55)

    db = SessionLocal()
    try:
        row = db.query(models.Student).filter_by(id=student.id).one()
        row.full_name = "Nombre Editado"
        row.course = "11B"
        db.commit()
    finally:
        db.close()

    first_page = client.get("/reports/")
    assert first_page.status_code == 200
    assert "55 registros" in first_page.text
    assert "Página 1 de 2" in first_page.text
    assert "Ana Perez" in first_page.text
    assert "Nombre Editado" not in first_page.text

    second_page = client.get("/reports/?page=2")
    assert second_page.status_code == 200
    assert "Página 2 de 2" in second_page.text

    clamped = client.get("/reports/?page=999")
    assert clamped.status_code == 200
    assert "Página 2 de 2" in clamped.text


def test_reports_reject_invalid_door_filter(client):
    login(client, "operator", USER_PASSWORD)

    assert client.get("/reports/?door_id=abc").status_code == 422
    assert (
        client.get(f"/reports/export?date_start={TODAY}&date_end={TODAY}&door_id=abc").status_code
        == 422
    )


def test_export_protects_against_formula_injection(client):
    login(client, "operator", USER_PASSWORD)
    student = create_student()
    _create_exit_logs(
        student.id,
        get_door_id(),
        student_name="=SUM(1)",
        course="+1",
        door_name="@SUM(1)",
    )

    response = client.get(f"/reports/export?date_start={TODAY}&date_end={TODAY}")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook["Reporte Salidas"]
    row = next(sheet.iter_rows(min_row=2, max_row=2))
    name_cell, course_cell, door_cell = row[2], row[3], row[4]

    assert name_cell.value == "'=SUM(1)"
    assert course_cell.value == "'+1"
    assert door_cell.value == "'@SUM(1)"
    assert all(cell.data_type != "f" for cell in row)


def test_reports_require_exit_role(client):
    login(client, "lunch", USER_PASSWORD)

    assert client.get("/reports/").status_code == 403
    assert client.get(f"/reports/export?date_start={TODAY}&date_end={TODAY}").status_code == 403
