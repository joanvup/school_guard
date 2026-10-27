import io
from datetime import datetime, time

from openpyxl import load_workbook

from app import auth, models
from app.database import SessionLocal
from app.utils.time import today_local
from tests.conftest import ADMIN_PASSWORD, USER_PASSWORD
from tests.helpers import create_student, csrf_headers, get_door_id, login

TODAY = today_local().isoformat()


def _create_exit_logs(student_id, door_id, count):
    db = SessionLocal()
    try:
        operator = db.query(models.User).filter_by(username="operator").one()
        for _ in range(count):
            db.add(
                models.ExitLog(
                    student_id=student_id,
                    operator_id=operator.id,
                    door_id=door_id,
                    timestamp=datetime.combine(today_local(), time(10, 0)),
                    student_name="Ana Perez",
                    course="10A",
                    door_name="Puerta Principal",
                    authorized_at_scan=True,
                )
            )
        db.commit()
    finally:
        db.close()


def test_dashboard_details_paginate_with_headers(client):
    login(client, "operator", USER_PASSWORD)
    student = create_student(student_id="S960")
    _create_exit_logs(student.id, get_door_id(), count=55)

    first = client.get(f"/api/dashboard/details?type=all&date={TODAY}&page=1&page_size=20")
    assert first.status_code == 200
    assert first.headers["X-Total-Count"] == "55"
    assert first.headers["X-Page"] == "1"
    assert first.headers["X-Page-Size"] == "20"
    assert len(first.json()) == 20

    third = client.get(f"/api/dashboard/details?type=all&date={TODAY}&page=3&page_size=20")
    assert len(third.json()) == 15

    beyond = client.get(f"/api/dashboard/details?type=all&date={TODAY}&page=4&page_size=20")
    assert beyond.json() == []

    invalid = client.get(f"/api/dashboard/details?type=all&date={TODAY}&page_size=500")
    assert invalid.status_code == 422


def test_lunch_export_sanitizes_formulas(client):
    login(client, "lunch", USER_PASSWORD)
    student = create_student(student_id="S961", full_name="=SUM(1)")

    created = client.post(
        "/lunch/process",
        json={"method": "qr", "code": auth.sign_qr_content(student.student_id)},
        headers=csrf_headers(client),
        follow_redirects=False,
    )
    assert created.status_code == 200

    response = client.get(f"/lunch/reports/export?date_start={TODAY}&date_end={TODAY}")
    assert response.status_code == 200

    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook["Almuerzos"]
    row = next(sheet.iter_rows(min_row=2, max_row=2))

    assert row[3].value == "'=SUM(1)"
    assert row[3].data_type != "f"
    assert row[4].value == "Estudiante"


def test_student_template_is_a_valid_workbook(client):
    login(client, "admin", ADMIN_PASSWORD)

    response = client.get("/students/template")

    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook["Plantilla"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert headers == ["ID", "Nombre Completo", "Curso", "Autorizado (SI/NO)"]
