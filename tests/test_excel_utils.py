import asyncio
import tempfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.utils import excel


def test_safe_cell_leaves_non_strings_untouched():
    assert excel.safe_cell(123) == 123
    assert excel.safe_cell(None) is None


def test_safe_cell_escapes_formula_prefixes_and_truncates():
    for prefix in excel.FORMULA_PREFIXES:
        assert excel.safe_cell(f"{prefix}SUM(A1)") == f"'{prefix}SUM(A1)"

    long_value = "x" * (excel.MAX_CELL_LENGTH + 10)
    assert len(excel.safe_cell(long_value)) == excel.MAX_CELL_LENGTH


def test_build_xlsx_escapes_formulas_and_preserves_non_strings():
    response = excel.build_xlsx(
        filename="reporte.xlsx",
        sheet_name="Datos",
        headers=["Nombre", "=Peligro"],
        rows=[
            ["=1+1", 5],
            ["Texto", None],
            ["+", "-"],
            ["@cmd", "\tTab"],
            ["\rCR", "y" * (excel.MAX_CELL_LENGTH + 5)],
        ],
    )

    assert response.media_type == excel.XLSX_MEDIA_TYPE
    assert response.filename == "reporte.xlsx"

    try:
        workbook = load_workbook(response.path, read_only=True)
        sheet = workbook["Datos"]
        rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
    finally:
        asyncio.run(response.background())

    assert not Path(response.path).exists()
    assert rows[0] == ("Nombre", "'=Peligro")
    assert rows[1] == ("'=1+1", 5)
    assert rows[2] == ("Texto",)
    assert rows[3] == ("'+", "'-")
    assert rows[4] == ("'@cmd", "'\tTab")
    assert rows[5][0] == "'\nCR"
    assert len(rows[5][1]) == excel.MAX_CELL_LENGTH


def test_build_xlsx_removes_temp_file_when_rendering_fails(monkeypatch):
    created = {}
    real_mkstemp = tempfile.mkstemp

    def tracking_mkstemp(*args, **kwargs):
        descriptor, path = real_mkstemp(*args, **kwargs)
        created["path"] = path
        return descriptor, path

    class FailingWorkbook:
        def __init__(self, *args, **kwargs):
            raise RuntimeError("fallo al generar el reporte")

    monkeypatch.setattr(excel.tempfile, "mkstemp", tracking_mkstemp)
    monkeypatch.setattr(excel, "Workbook", FailingWorkbook)

    with pytest.raises(RuntimeError, match="fallo al generar"):
        excel.build_xlsx(
            filename="roto.xlsx",
            sheet_name="Datos",
            headers=["Nombre"],
            rows=[["fila válida"]],
        )

    assert created["path"]
    assert not Path(created["path"]).exists()


def test_remove_file_ignores_missing_paths(tmp_path):
    excel._remove_file(str(tmp_path / "no-existe.xlsx"))
