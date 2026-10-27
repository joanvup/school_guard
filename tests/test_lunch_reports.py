import io
from datetime import datetime, timezone

from openpyxl import load_workbook

from app import models
from app.database import SessionLocal
from app.utils.time import today_local
from tests.conftest import USER_PASSWORD
from tests.helpers import csrf_headers, create_employee, create_student, login

TODAY = today_local().isoformat()


def _add_lunch_log(*, student_id=None, employee_id=None, delivered_type="Normal", person_name=None):
    db = SessionLocal()
    try:
        operator = db.query(models.User).filter_by(username="admin").one()
        db.add(
            models.LunchLog(
                service_date=today_local(),
                student_id=student_id,
                employee_id=employee_id,
                operator_id=operator.id,
                timestamp=datetime.now(timezone.utc),
                delivered_type=delivered_type,
                person_name=person_name,
            )
        )
        db.commit()
    finally:
        db.close()


def test_process_lunch_returns_controlled_500_on_unexpected_error(client, monkeypatch):
    login(client, "lunch", USER_PASSWORD)

    def boom(*args, **kwargs):
        raise RuntimeError("fallo inesperado")

    monkeypatch.setattr("app.routers.lunch.record_lunch", boom)

    response = client.post(
        "/lunch/process",
        json={"method": "qr", "code": "cualquiera"},
        headers=csrf_headers(client),
    )

    assert response.status_code == 500
    assert response.json() == {
        "status": "error",
        "message": "ERROR INTERNO AL REGISTRAR EL ALMUERZO",
    }


def test_missing_ticket_pages_return_404(client):
    login(client, "lunch", USER_PASSWORD)

    assert client.get("/lunch/tickets/no-existe").status_code == 404
    assert client.get("/lunch/tickets/no-existe/print").status_code == 404


def test_reports_filter_by_lunch_type_and_person_type(client):
    login(client, "lunch", USER_PASSWORD)
    student = create_student(student_id="S900", full_name="Alumno Filtro")
    employee = create_employee(doc_id="E900", full_name="Empleado Filtro")
    _add_lunch_log(student_id=student.id, delivered_type="Normal")
    _add_lunch_log(employee_id=employee.id, delivered_type="Especial")

    base = f"/lunch/reports?date_start={TODAY}&date_end={TODAY}"
    all_logs = client.get(base)
    assert all_logs.status_code == 200
    assert "Alumno Filtro" in all_logs.text
    assert "Empleado Filtro" in all_logs.text

    normal = client.get(f"{base}&lunch_type=Normal")
    assert normal.status_code == 200
    assert "Alumno Filtro" in normal.text
    assert "Empleado Filtro" not in normal.text

    todos = client.get(f"{base}&lunch_type=Todos")
    assert todos.status_code == 200
    assert "Alumno Filtro" in todos.text
    assert "Empleado Filtro" in todos.text

    only_students = client.get(f"{base}&person_type=student")
    assert only_students.status_code == 200
    assert "Alumno Filtro" in only_students.text
    assert "Empleado Filtro" not in only_students.text

    only_employees = client.get(f"{base}&person_type=employee")
    assert only_employees.status_code == 200
    assert "Empleado Filtro" in only_employees.text
    assert "Alumno Filtro" not in only_employees.text


def test_export_labels_employee_and_unknown_person_types(client):
    login(client, "lunch", USER_PASSWORD)
    employee = create_employee(doc_id="E901", full_name="Empleado Export")
    _add_lunch_log(employee_id=employee.id, person_name="Empleado Export")
    _add_lunch_log(person_name="Fantasma Sin Tipo")

    response = client.get(f"/lunch/reports/export?date_start={TODAY}&date_end={TODAY}")

    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    rows = {
        row[3].value: row[4].value
        for row in workbook["Almuerzos"].iter_rows(min_row=2)
        if row[3].value
    }
    assert rows["Empleado Export"] == "Empleado"
    assert rows["Fantasma Sin Tipo"] == "?"
